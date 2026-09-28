import os
from pathlib import Path
from typing import Dict, Any

def export_docx(title: str, text_content: str, output_path: str) -> str:
    """
    Exports structured advisory or executive summary to DOCX format.
    Uses python-docx if installed, with a markdown/text fallback.
    """
    path = Path(output_path)
    os.makedirs(path.parent, exist_ok=True)

    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        
        doc = docx.Document()
        
        # Add heading
        h = doc.add_heading(title, level=0)
        
        # Add paragraphs
        for line in text_content.splitlines():
            line_str = line.strip()
            if line_str.startswith("# "):
                doc.add_heading(line_str[2:], level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:], level=2)
            elif line_str.startswith("### "):
                doc.add_heading(line_str[4:], level=3)
            elif line_str.startswith("- ") or line_str.startswith("* "):
                doc.add_paragraph(line_str[2:], style='List Bullet')
            elif line_str:
                doc.add_paragraph(line_str)

        doc.save(str(path))
        return str(path)
    except Exception:
        # Fallback to plain document file
        fallback_path = path.with_suffix(".txt")
        with open(fallback_path, "w", encoding="utf-8") as f:
            f.write(f"{title}\n{'=' * len(title)}\n\n{text_content}")
        return str(fallback_path)
