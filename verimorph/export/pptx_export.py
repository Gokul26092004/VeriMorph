import os
from pathlib import Path
from typing import Dict, Any, List

def export_pptx(deck_data: Dict[str, Any], output_path: str) -> str:
    """
    Exports Presentation Deck to Microsoft PowerPoint (.pptx).
    Uses python-pptx if installed, with a markdown/presentation fallback.
    """
    path = Path(output_path)
    os.makedirs(path.parent, exist_ok=True)
    slides_data = deck_data.get("slides", [])

    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        prs = Presentation()
        # Set slide dimensions to widescreen 16:9
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        for s in slides_data:
            blank_slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_slide_layout)

            # Title
            tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(1.2))
            tf = tx_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = s.get("title", "Presentation Slide")
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(15, 23, 42)

            # Subtitle
            sub = s.get("subtitle", "")
            if sub:
                p2 = tf.add_paragraph()
                p2.text = sub
                p2.font.size = Pt(14)
                p2.font.color.rgb = RGBColor(100, 116, 139)

            # Bullets
            bullets = s.get("bullets", [])
            if bullets:
                content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(3.8))
                c_tf = content_box.text_frame
                c_tf.word_wrap = True
                for b_text in bullets:
                    bp = c_tf.add_paragraph()
                    bp.text = b_text.strip("• ")
                    bp.font.size = Pt(18)
                    bp.level = 0
                    bp.font.color.rgb = RGBColor(30, 41, 59)

            # Speaker Notes
            notes = s.get("speaker_notes", "")
            if notes and hasattr(slide, "notes_slide"):
                slide.notes_slide.notes_text_frame.text = notes

        prs.save(str(path))
        return str(path)
    except Exception:
        # Fallback to markdown slides
        fallback_path = path.with_suffix(".md")
        lines = [f"# {deck_data.get('title', 'Presentation Deck')}\n"]
        for s in slides_data:
            lines.append(f"## Slide {s.get('slide_number')}: {s.get('title')}")
            if s.get("subtitle"):
                lines.append(f"*{s.get('subtitle')}*\n")
            for b in s.get("bullets", []):
                lines.append(f"{b}")
            lines.append(f"\n> **Speaker Notes:** {s.get('speaker_notes')}\n")
            lines.append("---\n")
        with open(fallback_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return str(fallback_path)
